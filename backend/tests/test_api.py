from __future__ import annotations

import io
import json
import zipfile

import pytest

from app.config import get_settings


class TestHealthAndModels:
    def test_health_reports_ffmpeg(self, client) -> None:
        response = client.get("/api/health")
        assert response.status_code == 200
        body = response.json()
        assert body["api_key_configured"] is True
        assert body["base_url"].startswith("https://")
        assert "max_video_seconds" in body["limits"]

    def test_models_exposes_text_models(self, client) -> None:
        body = client.get("/api/models").json()
        assert body["default"] == get_settings().model_text_default
        assert [m["role"] for m in body["models"]] == ["text", "text", "text", "text", "omni"]

    def test_openapi_documents_every_route(self, client) -> None:
        paths = client.get("/openapi.json").json()["paths"]
        assert "/api/chat/stream" in paths
        assert "/api/files/analyze" in paths
        assert "/api/voice/speak" in paths


class TestFileValidation:
    def test_rejects_unsupported_type(self, client) -> None:
        response = client.post(
            "/api/files/analyze",
            files={"file": ("virus.exe", io.BytesIO(b"MZ"), "application/octet-stream")},
        )
        assert response.status_code == 415
        assert "no soportado" in response.json()["detail"]

    def test_rejects_oversized_image(self, client) -> None:
        big = io.BytesIO(b"\x89PNG\r\n\x1a\n" + b"0" * (21 * 1024 * 1024))
        response = client.post(
            "/api/files/analyze", files={"file": ("enorme.png", big, "image/png")}
        )
        assert response.status_code == 413
        assert "limite" in response.json()["detail"]

    def test_video_over_two_minutes_is_rejected(self, client, tmp_path) -> None:
        source = tmp_path / "largo.mp4"
        source.write_bytes(b"\x00" * 2048)
        _make_video_of_duration(source, seconds=200)

        response = client.post(
            "/api/files/analyze", files={"file": ("largo.mp4", source.open("rb"), "video/mp4")}
        )
        assert response.status_code == 400
        detail = response.json()["detail"]
        assert "2 minutos" in detail
        assert "Recorta el clip" in detail

    def test_file_is_removed_when_limits_fail(self, client, tmp_path) -> None:
        source = tmp_path / "largo2.mp4"
        source.write_bytes(b"\x00" * 2048)
        _make_video_of_duration(source, seconds=500)

        client.post(
            "/api/files/analyze", files={"file": ("largo2.mp4", source.open("rb"), "video/mp4")}
        )
        assert client.get("/api/files").json()["files"] == []


def _make_video_of_duration(path, seconds: int) -> None:
    """Crea un MP4 real de la duracion indicada usando FFmpeg."""
    import shutil
    import subprocess

    if shutil.which("ffmpeg") is None:
        pytest.skip("FFmpeg no disponible")
    subprocess.run(
        [
            "ffmpeg", "-y", "-f", "lavfi", "-i",
            f"color=c=blue:s=160x120:d={seconds}", "-f", "lavfi", "-i",
            "anullsrc=r=8000:cl=mono", "-shortest",
            "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac",
            str(path),
        ],
        check=True,
        capture_output=True,
    )


class TestStreamingChat:
    def test_requires_api_key(self, client, monkeypatch) -> None:
        from app.config import get_settings

        settings = get_settings()
        monkeypatch.setattr(settings, "nvidia_api_key", "")

        response = client.post(
            "/api/chat/stream",
            json={"messages": [{"role": "user", "content": "hola"}]},
        )
        assert response.status_code == 503
        assert "NVIDIA_API_KEY" in response.json()["detail"]

    def test_empty_conversation_yields_error_event(self, client) -> None:
        response = client.post("/api/chat/stream", json={"messages": []})
        assert response.status_code == 200
        assert "event: error" in response.text

    def test_sse_frame_shape(self, client) -> None:
        """Comprueba que el serializador SSE produce frames parseables."""
        from app.api.routes_chat import sse

        frame = sse("delta", {"delta": "hola\nmundo"})
        assert frame.startswith("event: delta\n")
        payload = json.loads(frame.split("data: ", 1)[1].strip())
        assert payload == {"delta": "hola\nmundo"}
        assert frame.endswith("\n\n")


class TestKnowledgeSearch:
    def test_returns_503_when_rag_disabled(self, client) -> None:
        response = client.get("/api/knowledge/search?q=hola")
        assert response.status_code == 503
        assert "RAG" in response.json()["detail"]

    def test_rejects_short_query(self, client) -> None:
        assert client.get("/api/knowledge/search?q=a").status_code == 422


class TestVoiceRoutes:
    def test_speak_requires_text(self, client) -> None:
        assert client.post("/api/voice/speak", json={"text": ""}).status_code == 422

    def test_transcribe_rejects_missing_file(self, client) -> None:
        assert client.post("/api/voice/transcribe").status_code == 422


class TestAnalyzeFlow:
    """Verifica el flujo completo de subida con la llamada a NVIDIA simulada."""

    @staticmethod
    def _stub(client, reply: str = "Analisis simulado."):
        from app.main import app

        calls: list[dict] = []

        async def fake_complete(model, messages, **kwargs):
            calls.append({"model": model, "messages": messages, **kwargs})
            return "razonamiento simulado", reply

        client.app.state.ctx.client.complete = fake_complete
        return calls

    def _png(self) -> io.BytesIO:
        from PIL import Image

        buffer = io.BytesIO()
        Image.new("RGB", (64, 48), (12, 140, 60)).save(buffer, format="PNG")
        buffer.seek(0)
        return buffer

    def test_image_is_analyzed_and_registered(self, client) -> None:
        calls = self._stub(client)
        response = client.post(
            "/api/files/analyze", files={"file": ("foto.png", self._png(), "image/png")}
        )
        assert response.status_code == 200
        body = response.json()
        assert body["analysis"] == "Analisis simulado."
        assert body["attachment"]["kind"] == "image"
        assert body["attachment"]["width"] == 64
        assert body["attachment"]["height"] == 48

        listed = client.get("/api/files").json()["files"]
        assert [item["name"] for item in listed] == ["foto.png"]

        raw = client.get(body["attachment"]["url"])
        assert raw.status_code == 200
        assert raw.headers["content-type"] == "image/png"

    def test_image_payload_uses_omni_with_data_url(self, client) -> None:
        calls = self._stub(client)
        client.post("/api/files/analyze", files={"file": ("foto.png", self._png(), "image/png")})

        assert len(calls) == 1
        call = calls[0]
        assert call["model"] == "nvidia/nemotron-3-nano-omni-30b-a3b-reasoning"
        content = call["messages"][0]["content"]
        assert content[0]["type"] == "text"
        assert content[1]["type"] == "image_url"
        assert content[1]["image_url"]["url"].startswith("data:image/png;base64,")
        # El prompt de analisis va en ingles pero exige respuesta en espanol.
        assert "español" in content[0]["text"].lower()
        assert call["thinking"] is False

    def test_custom_prompt_enables_reasoning(self, client) -> None:
        calls = self._stub(client)
        client.post(
            "/api/files/analyze",
            files={"file": ("foto.png", self._png(), "image/png")},
            data={"prompt": "Cuantas personas hay?"},
        )
        assert calls[0]["messages"][0]["content"][0]["text"].startswith("Cuantas personas hay?")
        assert calls[0]["thinking"] is True

    def test_nvidia_failure_cleans_up_and_returns_502(self, client) -> None:
        from app.main import app

        async def failing(model, messages, **kwargs):
            from app.services.nvidia_client import NvidiaError

            raise NvidiaError("cuota agotada")

        client.app.state.ctx.client.complete = failing

        response = client.post(
            "/api/files/analyze", files={"file": ("foto.png", self._png(), "image/png")}
        )
        assert response.status_code == 502
        assert "cuota agotada" in response.json()["detail"]
        assert client.get("/api/files").json()["files"] == []

        uploads = (client.app.state.ctx.settings.uploads_dir)
        assert list(uploads.glob("*")) == []

    def test_text_document_is_analyzed(self, client) -> None:
        calls = self._stub(client, reply="Resumen del documento.")
        response = client.post(
            "/api/files/analyze",
            files={"file": ("notas.txt", io.BytesIO("contenido plano".encode()), "text/plain")},
        )
        assert response.status_code == 200
        assert response.json()["attachment"]["kind"] == "document"
        assert calls[0]["model"] == "nvidia/nemotron-3-nano-omni-30b-a3b-reasoning"
        assert "contenido plano" in calls[0]["messages"][0]["content"][1]["text"]

    def test_delete_removes_file_and_row(self, client) -> None:
        self._stub(client)
        body = client.post(
            "/api/files/analyze", files={"file": ("foto.png", self._png(), "image/png")}
        ).json()
        file_id = body["attachment"]["file_id"]

        assert client.delete(f"/api/files/{file_id}").json() == {"deleted": True}
        assert client.get("/api/files").json()["files"] == []
        assert client.get(f"/api/files/{file_id}/raw").status_code == 404


class TestDocumentHelpers:
    def test_extracts_text_from_docx(self, tmp_path) -> None:
        import docx

        document = docx.Document()
        document.add_paragraph("Titulo del informe")
        document.add_paragraph("Contenido relevante")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Columna A"
        table.rows[0].cells[1].text = "Columna B"
        path = tmp_path / "informe.docx"
        document.save(path)

        from app.services.media import extract_docx

        text = extract_docx(path)
        assert "Titulo del informe" in text
        assert "Columna A | Columna B" in text

    def test_reads_plain_text(self, tmp_path) -> None:
        from app.services.media import extract_plain_text

        path = tmp_path / "notas.md"
        path.write_text("# Notas\ncon acentos: ñ, á, é", encoding="utf-8")
        assert "con acentos" in extract_plain_text(path)

    def test_extracts_pdf_text(self, tmp_path) -> None:
        fitz = pytest.importorskip("fitz")

        path = tmp_path / "factura.pdf"
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 100), "Total: 1234.56 EUR")
        document.save(path)
        document.close()

        from app.services.media import extract_pdf

        text, page_count, images = extract_pdf(path, max_pages=5)
        assert page_count == 1
        assert "1234.56" in text
        assert len(images) == 1

    def test_docx_roundtrip_is_valid_zip(self, tmp_path) -> None:
        path = tmp_path / "x.docx"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
        assert zipfile.is_zipfile(path)
